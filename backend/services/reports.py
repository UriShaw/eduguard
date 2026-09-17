"""Xuất báo cáo Excel và nhập dữ liệu hàng loạt từ file CSV/Excel."""
from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from models import AcademicResult, Attendance, LearningInteraction, Student
from services import latest_for, ml

# ===========================================================================
# Báo cáo Excel
# ===========================================================================

REPORT_NOTE = ("Chuyên cần và mức nguy cơ là số liệu GẦN NHẤT của sinh viên tại thời điểm xuất "
               "báo cáo, không tách riêng theo học kỳ đang lọc.")


def semesters(db: Session, scope: int | None) -> list[str]:
    query = select(AcademicResult.semester).distinct()
    if scope is not None:
        query = query.join(Student).where(Student.advisor_user_id == scope)
    return sorted(db.scalars(query), reverse=True)


def report_workbook(db: Session, scope: int | None, class_name: str | None, semester: str | None) -> BytesIO:
    """
    Hai sheet: Tổng hợp theo lớp và học kỳ, Chi tiết từng sinh viên.

    Giới hạn của mô hình dữ liệu được in thẳng vào file: chỉ bảng kết quả học
    tập có cột học kỳ, nên chuyên cần và mức nguy cơ là số gần nhất — người
    nhận file cần biết điều này để không đọc nhầm.
    """
    query = select(AcademicResult, Student).join(Student)
    if scope is not None:
        query = query.where(Student.advisor_user_id == scope)
    if class_name:
        query = query.where(Student.class_name == class_name)
    if semester:
        query = query.where(AcademicResult.semester == semester)
    pairs = db.execute(query.order_by(Student.class_name, AcademicResult.semester, Student.full_name)).all()

    latest = latest_for(db, {s.student_id for _, s in pairs})
    rows = []
    for record, student in pairs:
        att, pred = latest["attendance"].get(student.student_id), latest["prediction"].get(student.student_id)
        rows.append([student.class_name or "—", record.semester, student.student_code, student.full_name,
                     record.gpa, record.failed_subjects, att.attendance_rate if att else None,
                     pred.risk_level if pred else "Chưa dự đoán"])

    book = Workbook()
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1D4ED8")

    def header(sheet, row: int, titles: list[str], widths: list[int]):
        for col, (title, width) in enumerate(zip(titles, widths), start=1):
            cell = sheet.cell(row=row, column=col, value=title)
            cell.font, cell.fill, cell.alignment = header_font, header_fill, Alignment(horizontal="center")
            sheet.column_dimensions[cell.column_letter].width = width

    summary_sheet = book.active
    summary_sheet.title = "Tổng hợp"
    summary_sheet["A1"] = "BÁO CÁO NGUY CƠ BỎ HỌC THEO LỚP VÀ HỌC KỲ"
    summary_sheet["A1"].font = Font(bold=True, size=14, color="1D4ED8")
    summary_sheet["A2"] = REPORT_NOTE
    summary_sheet["A2"].font = Font(italic=True, size=9, color="666666")
    header(summary_sheet, 4, ["Lớp", "Học kỳ", "Số SV", "GPA TB", "Chuyên cần TB (%)",
                              "Nguy cơ thấp", "Trung bình", "Cao + Rất cao"], [16, 10, 8, 10, 18, 13, 12, 14])

    groups: dict = {}
    for row in rows:
        groups.setdefault((row[0], row[1]), []).append(row)
    for line, ((cls, sem), items) in enumerate(sorted(groups.items()), start=5):
        rates = [i[6] for i in items if i[6] is not None]
        summary_sheet.append([cls, sem, len(items), round(sum(i[4] for i in items) / len(items), 2),
                              round(sum(rates) / len(rates), 1) if rates else None,
                              sum(i[7] == "Thấp" for i in items), sum(i[7] == "Trung bình" for i in items),
                              sum(i[7] in ml.AT_RISK_LEVELS for i in items)])

    detail = book.create_sheet("Chi tiết")
    header(detail, 1, ["Lớp", "Học kỳ", "Mã SV", "Họ và tên", "GPA", "Môn trượt",
                       "Chuyên cần gần nhất (%)", "Mức nguy cơ gần nhất"], [16, 10, 14, 26, 8, 10, 22, 20])
    for row in rows:
        detail.append(row)
    detail.freeze_panes = "A2"

    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


# ===========================================================================
# Nhập dữ liệu từ file
# ===========================================================================
# Hai nguyên tắc:
#  1. Không dừng ở dòng lỗi đầu tiên — file thật luôn có vài dòng sai, dừng
#     ngay buộc người dùng sửa từng dòng một qua nhiều lượt tải lên.
#  2. Kiểm tra trọn file trước, ghi sau, trong MỘT giao dịch — hỏng giữa chừng
#     thì không để CSDL ở trạng thái nhập được một nửa.

STUDENT_COLUMNS = ["student_code", "full_name", "email", "phone", "class_name", "major",
                   "enrollment_year", "status"]
METRIC_COLUMNS = ["student_code", "semester", "gpa", "failed_subjects", "credits_completed",
                  "credits_registered", "attendance_rate", "total_sessions", "absent_sessions",
                  "login_count", "assignment_submitted", "assignment_missing", "forum_posts",
                  "video_views", "learning_hours"]
ROW_OFFSET = 2  # dòng 1 là tiêu đề, pandas đếm từ 0 — để số dòng báo lỗi khớp Excel


class ImportFileError(Exception):
    """Lỗi cả file (sai định dạng, thiếu cột) — khác với lỗi của từng dòng."""


def read_upload(filename: str, content: bytes) -> pd.DataFrame:
    name = filename.lower()
    if name.endswith(".csv"):
        frame = pd.read_csv(BytesIO(content))
    elif name.endswith((".xlsx", ".xls")):
        frame = pd.read_excel(BytesIO(content))
    else:
        raise ImportFileError("Chỉ hỗ trợ file .csv, .xlsx hoặc .xls")
    frame.columns = [str(c).strip() for c in frame.columns]
    return frame


def _text(value) -> str | None:
    return None if pd.isna(value) or not str(value).strip() else str(value).strip()


def _commit(db: Session, pending: list, errors: list[dict]) -> int:
    if not pending:
        return 0
    try:
        db.add_all(pending)
        db.commit()
    except Exception as exc:
        db.rollback()
        # Cả lô bị huỷ — nói rõ là KHÔNG dòng nào vào, đừng để người dùng đoán.
        errors.append({"row": None, "message": f"Không ghi được dữ liệu, đã huỷ toàn bộ: {exc}"})
        return 0
    return len(pending)


def import_students(db: Session, frame: pd.DataFrame) -> dict:
    """
    Thêm sinh viên mới. Mã đã tồn tại BỊ BỎ QUA, không ghi đè: một file cũ tải
    nhầm không được phép xoá sạch thông tin đã cập nhật của hàng trăm người.
    """
    missing = {"student_code", "full_name"} - set(frame.columns)
    if missing:
        raise ImportFileError(f"File thiếu cột bắt buộc: {', '.join(sorted(missing))}")

    existing = set(db.scalars(select(Student.student_code)))
    seen, pending, errors = set(), [], []

    for index, row in frame.iterrows():
        line = index + ROW_OFFSET
        code, name = _text(row.get("student_code")), _text(row.get("full_name"))
        if not code or not name:
            errors.append({"row": line, "message": "Thiếu mã sinh viên hoặc họ tên"})
        elif code in existing:
            errors.append({"row": line, "message": f"Mã '{code}' đã có trong hệ thống — bỏ qua"})
        elif code in seen:
            # Bắt ở đây, nếu không ràng buộc UNIQUE làm hỏng cả lô ở bước ghi.
            errors.append({"row": line, "message": f"Mã '{code}' bị lặp trong file"})
        else:
            seen.add(code)
            status = _text(row.get("status"))
            year = row.get("enrollment_year")
            pending.append(Student(
                student_code=code, full_name=name, email=_text(row.get("email")), phone=_text(row.get("phone")),
                class_name=_text(row.get("class_name")), major=_text(row.get("major")),
                enrollment_year=int(year) if year is not None and not pd.isna(year) else None,
                status=status if status in Student.__table__.c.status.type.enums else "active"))

    return {"imported": _commit(db, pending, errors), "errors": errors}


def import_metrics(db: Session, frame: pd.DataFrame, user) -> dict:
    """
    Chỉ số cho nhiều sinh viên, khớp theo mã. Giảng viên nhập nhầm sinh viên
    không thuộc mình phụ trách sẽ thấy lỗi từng dòng, thay vì bị bỏ qua âm thầm
    rồi tưởng dữ liệu đã vào.
    """
    missing = {"student_code", "semester", "gpa", "attendance_rate"} - set(frame.columns)
    if missing:
        raise ImportFileError(f"File thiếu cột bắt buộc: {', '.join(sorted(missing))}")

    by_code = {s.student_code: s for s in db.scalars(select(Student).where(
        Student.student_code.in_([str(c).strip() for c in frame["student_code"].dropna()])))}
    today, pending, errors, imported_rows = date.today(), [], [], 0

    def number(row, key, cast=int):
        value = row.get(key)
        return cast(0) if value is None or pd.isna(value) else cast(value)

    for index, row in frame.iterrows():
        line = index + ROW_OFFSET
        student = by_code.get(_text(row.get("student_code")) or "")
        try:
            gpa, rate = float(row["gpa"]), float(row["attendance_rate"])
        except (TypeError, ValueError):
            errors.append({"row": line, "message": "GPA hoặc chuyên cần không phải số"})
            continue

        if student is None:
            errors.append({"row": line, "message": f"Không tìm thấy sinh viên '{_text(row.get('student_code'))}'"})
        elif not user.can_edit(student):
            errors.append({"row": line, "message": f"Bạn không phụ trách sinh viên '{student.student_code}'"})
        elif not _text(row.get("semester")):
            errors.append({"row": line, "message": "Thiếu học kỳ"})
        elif not (0 <= gpa <= 10 and 0 <= rate <= 100):
            errors.append({"row": line, "message": f"GPA {gpa} hoặc chuyên cần {rate} ngoài khoảng cho phép"})
        else:
            sid = student.student_id
            pending += [
                AcademicResult(student_id=sid, semester=_text(row["semester"]), gpa=gpa,
                               failed_subjects=number(row, "failed_subjects"),
                               credits_completed=number(row, "credits_completed"),
                               credits_registered=number(row, "credits_registered")),
                Attendance(student_id=sid, period_start=today, period_end=today, attendance_rate=rate,
                           total_sessions=number(row, "total_sessions"),
                           absent_sessions=number(row, "absent_sessions")),
                LearningInteraction(student_id=sid, period_start=today, period_end=today,
                                    login_count=number(row, "login_count"),
                                    assignment_submitted=number(row, "assignment_submitted"),
                                    assignment_missing=number(row, "assignment_missing"),
                                    forum_posts=number(row, "forum_posts"), video_views=number(row, "video_views"),
                                    learning_hours=number(row, "learning_hours", float)),
            ]
            imported_rows += 1

    return {"imported": imported_rows if _commit(db, pending, errors) else 0, "errors": errors}


def template_csv(kind: str) -> tuple[bytes, str]:
    samples = {
        "students": (STUDENT_COLUMNS, ["SV20250001", "Nguyễn Văn An", "an@example.edu.vn", "0901234567",
                                       "CNTT-K25-01", "Công nghệ thông tin", "2025", "active"],
                     "mau_sinh_vien.csv"),
        "metrics": (METRIC_COLUMNS, ["SV20250001", "2025.1", "7.5", "0", "68", "18", "86.5", "45", "6",
                                     "24", "11", "1", "3", "28", "12.5"], "mau_chi_so.csv"),
    }
    if kind not in samples:
        raise ImportFileError(f"Không có file mẫu '{kind}'")
    columns, sample, filename = samples[kind]
    # BOM để Excel trên Windows mở đúng tiếng Việt.
    return (",".join(columns) + "\n" + ",".join(sample) + "\n").encode("utf-8-sig"), filename
