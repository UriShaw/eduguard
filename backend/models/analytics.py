"""
Tầng Model — thống kê, cảnh báo sớm, báo cáo Excel, nhập dữ liệu từ file.

Mọi con số tính trên LẦN DỰ ĐOÁN GẦN NHẤT của mỗi sinh viên: một người được
chạy lại năm lần không vì thế mà đáng lo gấp năm. Việc đếm đẩy xuống SQL thay
vì nạp bản ghi lên rồi cộng bằng vòng lặp Python.
"""
from datetime import date, timedelta
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import and_, func, select
from sqlalchemy.orm import Session

from models import (AcademicResult, Attendance, Intervention, LearningInteraction,
                    Prediction, Student, ml)
from models.services import latest_for, newest_prediction_subquery

# Tăng từ ngưỡng này (điểm phần trăm) giữa hai lần dự đoán liên tiếp thì coi là
# "đang xấu đi nhanh". 15 điểm đủ lớn để không báo động vì dao động thông thường.
DETERIORATION_THRESHOLD = 15.0


def _newest(scope: int | None):
    """Truy vấn trả đúng bản ghi dự đoán mới nhất của mỗi sinh viên trong phạm vi."""
    newest = newest_prediction_subquery()
    query = (select(Prediction)
             .join(newest, and_(Prediction.student_id == newest.c.sid, Prediction.created_at == newest.c.at)))
    if scope is not None:
        query = query.join(Student, Student.student_id == Prediction.student_id).where(
            Student.advisor_user_id == scope)
    return query


# ===========================================================================
# Dashboard
# ===========================================================================

def summary(db: Session, scope: int | None) -> dict:
    students = select(func.count(Student.student_id))
    if scope is not None:
        students = students.where(Student.advisor_user_id == scope)
    total = db.scalar(students) or 0

    newest = _newest(scope).subquery()
    counts = dict(db.execute(select(newest.c.risk_level, func.count()).group_by(newest.c.risk_level)).all())

    # Luôn đủ bốn khoá kể cả mức rỗng — thiếu khoá thì biểu đồ nhảy cột.
    distribution = {level: int(counts.get(level, 0)) for level in ml.RISK_LEVELS}
    predicted = sum(distribution.values())
    return {
        "total_students": total,
        "predicted": predicted,
        "not_predicted": total - predicted,
        "distribution": distribution,
        "at_risk": sum(distribution[level] for level in ml.AT_RISK_LEVELS),
    }


def scatter(db: Session, scope: int | None, limit: int = 500) -> list[dict]:
    """GPA và chuyên cần đối chiếu xác suất. Giới hạn điểm: vài nghìn chấm vừa nặng vừa không đọc được."""
    predictions = db.scalars(_newest(scope).limit(limit)).all()
    latest = latest_for(db, (p.student_id for p in predictions))

    points = []
    for p in predictions:
        academic, attendance = latest["academic"].get(p.student_id), latest["attendance"].get(p.student_id)
        if academic and attendance:
            points.append({"gpa": academic.gpa, "attendance": attendance.attendance_rate,
                           "risk": round(p.probability * 100, 1), "level": p.risk_level})
    return points


def by_class(db: Session, scope: int | None, top: int = 8) -> list[dict]:
    newest = _newest(scope).subquery()
    rows = db.execute(
        select(Student.class_name, func.count())
        .join(newest, newest.c.student_id == Student.student_id)
        .where(newest.c.risk_level.in_(ml.AT_RISK_LEVELS))
        .group_by(Student.class_name).order_by(func.count().desc()).limit(top)).all()
    return [{"class_name": name or "Chưa phân lớp", "count": int(count)} for name, count in rows]


# ===========================================================================
# CẢNH BÁO SỚM
# ===========================================================================
# Dashboard trả lời "ai đang có nguy cơ". Cảnh báo sớm trả lời câu hỏi cố vấn
# thực sự cần hằng ngày: "hôm nay tôi nên xử lý ai trước". Ba danh sách,
# mỗi danh sách là một loại việc bị bỏ lỡ khác nhau.

def alerts(db: Session, scope: int | None) -> dict:
    newest = {p.student_id: p for p in db.scalars(_newest(scope))}

    # 1. Đang xấu đi nhanh: so lần dự đoán mới nhất với lần liền trước.
    previous_q = (select(Prediction)
                  .where(Prediction.student_id.in_(list(newest)))
                  .order_by(Prediction.student_id, Prediction.created_at.desc()))
    previous: dict[int, Prediction] = {}
    for p in db.scalars(previous_q):
        if p.prediction_id != newest[p.student_id].prediction_id and p.student_id not in previous:
            previous[p.student_id] = p

    worsening = []
    for sid, last in newest.items():
        before = previous.get(sid)
        if before:
            delta = round((last.probability - before.probability) * 100, 1)
            if delta >= DETERIORATION_THRESHOLD:
                worsening.append((sid, {"delta": delta, "from": before.probability, "to": last.probability}))
    worsening.sort(key=lambda item: item[1]["delta"], reverse=True)

    # 2. Nguy cơ cao mà CHƯA có việc can thiệp nào đang mở — dễ lọt nhất.
    at_risk = [sid for sid, p in newest.items() if p.risk_level in ml.AT_RISK_LEVELS]
    has_plan = set(db.scalars(select(Intervention.student_id).where(
        Intervention.student_id.in_(at_risk), Intervention.status != "completed")))
    unplanned = sorted((sid for sid in at_risk if sid not in has_plan),
                       key=lambda sid: newest[sid].probability, reverse=True)

    # 3. Việc can thiệp đã quá hạn.
    overdue_q = select(Intervention).where(Intervention.status != "completed",
                                           Intervention.due_date < date.today())
    if scope is not None:
        overdue_q = overdue_q.join(Student).where(Student.advisor_user_id == scope)
    overdue = db.scalars(overdue_q.order_by(Intervention.due_date)).all()

    involved = {sid for sid, _ in worsening} | set(unplanned) | {i.student_id for i in overdue}
    students = {s.student_id: s for s in db.scalars(select(Student).where(Student.student_id.in_(list(involved))))}

    def card(sid: int, **extra) -> dict:
        s, p = students[sid], newest.get(sid)
        return {"student_id": sid, "student_code": s.student_code, "full_name": s.full_name,
                "class_name": s.class_name, "risk_level": p.risk_level if p else None,
                "probability": p.probability if p else None, **extra}

    return {
        "worsening": [card(sid, **info) for sid, info in worsening],
        "unplanned": [card(sid) for sid in unplanned],
        "overdue": [card(i.student_id, intervention_id=i.intervention_id, title=i.title,
                         due_date=i.due_date.isoformat(), days_overdue=(date.today() - i.due_date).days)
                    for i in overdue],
    }


def alert_count(db: Session, scope: int | None) -> int:
    data = alerts(db, scope)
    return len(data["worsening"]) + len(data["unplanned"]) + len(data["overdue"])


# ===========================================================================
# Báo cáo Excel
# ===========================================================================

REPORT_NOTE = ("Chuyên cần và mức nguy cơ là số liệu GẦN NHẤT của sinh viên tại thời điểm xuất "
               "báo cáo, không tách riêng theo học kỳ đang lọc.")


def semesters(db: Session, scope: int | None) -> list[str]:
    query = select(AcademicResult.semester).distinct()
    if scope is not None:
        query = query.join(Student).where(Student.advisor_user_id == scope)
    return sorted(db.scalars(query), reverse=True)


def report_workbook(db: Session, scope: int | None, class_name: str | None, semester: str | None) -> BytesIO:
    """
    Hai sheet: Tổng hợp theo lớp và học kỳ, Chi tiết từng sinh viên.

    Giới hạn của mô hình dữ liệu được in thẳng vào file: chỉ bảng kết quả học
    tập có cột học kỳ, nên chuyên cần và mức nguy cơ là số gần nhất — người
    nhận file cần biết điều này để không đọc nhầm.
    """
    query = select(AcademicResult, Student).join(Student)
    if scope is not None:
        query = query.where(Student.advisor_user_id == scope)
    if class_name:
        query = query.where(Student.class_name == class_name)
    if semester:
        query = query.where(AcademicResult.semester == semester)
    pairs = db.execute(query.order_by(Student.class_name, AcademicResult.semester, Student.full_name)).all()

    latest = latest_for(db, {s.student_id for _, s in pairs})
    rows = []
    for record, student in pairs:
        att, pred = latest["attendance"].get(student.student_id), latest["prediction"].get(student.student_id)
        rows.append([student.class_name or "—", record.semester, student.student_code, student.full_name,
                     record.gpa, record.failed_subjects, att.attendance_rate if att else None,
                     pred.risk_level if pred else "Chưa dự đoán"])

    book = Workbook()
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1D4ED8")

    def header(sheet, row: int, titles: list[str], widths: list[int]):
        for col, (title, width) in enumerate(zip(titles, widths), start=1):
            cell = sheet.cell(row=row, column=col, value=title)
            cell.font, cell.fill, cell.alignment = header_font, header_fill, Alignment(horizontal="center")
            sheet.column_dimensions[cell.column_letter].width = width

    summary_sheet = book.active
    summary_sheet.title = "Tổng hợp"
    summary_sheet["A1"] = "BÁO CÁO NGUY CƠ BỎ HỌC THEO LỚP VÀ HỌC KỲ"
    summary_sheet["A1"].font = Font(bold=True, size=14, color="1D4ED8")
    summary_sheet["A2"] = REPORT_NOTE
    summary_sheet["A2"].font = Font(italic=True, size=9, color="666666")
    header(summary_sheet, 4, ["Lớp", "Học kỳ", "Số SV", "GPA TB", "Chuyên cần TB (%)",
                              "Nguy cơ thấp", "Trung bình", "Cao + Rất cao"], [16, 10, 8, 10, 18, 13, 12, 14])

    groups: dict = {}
    for row in rows:
        groups.setdefault((row[0], row[1]), []).append(row)
    for line, ((cls, sem), items) in enumerate(sorted(groups.items()), start=5):
        rates = [i[6] for i in items if i[6] is not None]
        summary_sheet.append([cls, sem, len(items), round(sum(i[4] for i in items) / len(items), 2),
                              round(sum(rates) / len(rates), 1) if rates else None,
                              sum(i[7] == "Thấp" for i in items), sum(i[7] == "Trung bình" for i in items),
                              sum(i[7] in ml.AT_RISK_LEVELS for i in items)])

    detail = book.create_sheet("Chi tiết")
    header(detail, 1, ["Lớp", "Học kỳ", "Mã SV", "Họ và tên", "GPA", "Môn trượt",
                       "Chuyên cần gần nhất (%)", "Mức nguy cơ gần nhất"], [16, 10, 14, 26, 8, 10, 22, 20])
    for row in rows:
        detail.append(row)
    detail.freeze_panes = "A2"

    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


# ===========================================================================
# Nhập dữ liệu từ file
# ===========================================================================
# Hai nguyên tắc:
#  1. Không dừng ở dòng lỗi đầu tiên — file thật luôn có vài dòng sai, dừng
#     ngay buộc người dùng sửa từng dòng một qua nhiều lượt tải lên.
#  2. Kiểm tra trọn file trước, ghi sau, trong MỘT giao dịch — hỏng giữa chừng
#     thì không để CSDL ở trạng thái nhập được một nửa.

STUDENT_COLUMNS = ["student_code", "full_name", "email", "phone", "class_name", "major",
                   "enrollment_year", "status"]
METRIC_COLUMNS = ["student_code", "semester", "gpa", "failed_subjects", "credits_completed",
                  "credits_registered", "attendance_rate", "total_sessions", "absent_sessions",
                  "login_count", "assignment_submitted", "assignment_missing", "forum_posts",
                  "video_views", "learning_hours"]
ROW_OFFSET = 2  # dòng 1 là tiêu đề, pandas đếm từ 0 — để số dòng báo lỗi khớp Excel


class ImportFileError(Exception):
    """Lỗi cả file (sai định dạng, thiếu cột) — khác với lỗi của từng dòng."""


def read_upload(filename: str, content: bytes) -> pd.DataFrame:
    name = filename.lower()
    if name.endswith(".csv"):
        frame = pd.read_csv(BytesIO(content))
    elif name.endswith((".xlsx", ".xls")):
        frame = pd.read_excel(BytesIO(content))
    else:
        raise ImportFileError("Chỉ hỗ trợ file .csv, .xlsx hoặc .xls")
    frame.columns = [str(c).strip() for c in frame.columns]
    return frame


def _text(value) -> str | None:
    return None if pd.isna(value) or not str(value).strip() else str(value).strip()


def _commit(db: Session, pending: list, errors: list[dict]) -> int:
    if not pending:
        return 0
    try:
        db.add_all(pending)
        db.commit()
    except Exception as exc:
        db.rollback()
        # Cả lô bị huỷ — nói rõ là KHÔNG dòng nào vào, đừng để người dùng đoán.
        errors.append({"row": None, "message": f"Không ghi được dữ liệu, đã huỷ toàn bộ: {exc}"})
        return 0
    return len(pending)


def import_students(db: Session, frame: pd.DataFrame) -> dict:
    """
    Thêm sinh viên mới. Mã đã tồn tại BỊ BỎ QUA, không ghi đè: một file cũ tải
    nhầm không được phép xoá sạch thông tin đã cập nhật của hàng trăm người.
    """
    missing = {"student_code", "full_name"} - set(frame.columns)
    if missing:
        raise ImportFileError(f"File thiếu cột bắt buộc: {', '.join(sorted(missing))}")

    existing = set(db.scalars(select(Student.student_code)))
    seen, pending, errors = set(), [], []

    for index, row in frame.iterrows():
        line = index + ROW_OFFSET
        code, name = _text(row.get("student_code")), _text(row.get("full_name"))
        if not code or not name:
            errors.append({"row": line, "message": "Thiếu mã sinh viên hoặc họ tên"})
        elif code in existing:
            errors.append({"row": line, "message": f"Mã '{code}' đã có trong hệ thống — bỏ qua"})
        elif code in seen:
            # Bắt ở đây, nếu không ràng buộc UNIQUE làm hỏng cả lô ở bước ghi.
            errors.append({"row": line, "message": f"Mã '{code}' bị lặp trong file"})
        else:
            seen.add(code)
            status = _text(row.get("status"))
            year = row.get("enrollment_year")
            pending.append(Student(
                student_code=code, full_name=name, email=_text(row.get("email")), phone=_text(row.get("phone")),
                class_name=_text(row.get("class_name")), major=_text(row.get("major")),
                enrollment_year=int(year) if year is not None and not pd.isna(year) else None,
                status=status if status in Student.__table__.c.status.type.enums else "active"))

    return {"imported": _commit(db, pending, errors), "errors": errors}


def import_metrics(db: Session, frame: pd.DataFrame, user) -> dict:
    """
    Chỉ số cho nhiều sinh viên, khớp theo mã. Giảng viên nhập nhầm sinh viên
    không thuộc mình phụ trách sẽ thấy lỗi từng dòng, thay vì bị bỏ qua âm thầm
    rồi tưởng dữ liệu đã vào.
    """
    missing = {"student_code", "semester", "gpa", "attendance_rate"} - set(frame.columns)
    if missing:
        raise ImportFileError(f"File thiếu cột bắt buộc: {', '.join(sorted(missing))}")

    by_code = {s.student_code: s for s in db.scalars(select(Student).where(
        Student.student_code.in_([str(c).strip() for c in frame["student_code"].dropna()])))}
    today, pending, errors, imported_rows = date.today(), [], [], 0

    def number(row, key, cast=int):
        value = row.get(key)
        return cast(0) if value is None or pd.isna(value) else cast(value)

    for index, row in frame.iterrows():
        line = index + ROW_OFFSET
        student = by_code.get(_text(row.get("student_code")) or "")
        try:
            gpa, rate = float(row["gpa"]), float(row["attendance_rate"])
        except (TypeError, ValueError):
            errors.append({"row": line, "message": "GPA hoặc chuyên cần không phải số"})
            continue

        if student is None:
            errors.append({"row": line, "message": f"Không tìm thấy sinh viên '{_text(row.get('student_code'))}'"})
        elif not user.can_edit(student):
            errors.append({"row": line, "message": f"Bạn không phụ trách sinh viên '{student.student_code}'"})
        elif not _text(row.get("semester")):
            errors.append({"row": line, "message": "Thiếu học kỳ"})
        elif not (0 <= gpa <= 10 and 0 <= rate <= 100):
            errors.append({"row": line, "message": f"GPA {gpa} hoặc chuyên cần {rate} ngoài khoảng cho phép"})
        else:
            sid = student.student_id
            pending += [
                AcademicResult(student_id=sid, semester=_text(row["semester"]), gpa=gpa,
                               failed_subjects=number(row, "failed_subjects"),
                               credits_completed=number(row, "credits_completed"),
                               credits_registered=number(row, "credits_registered")),
                Attendance(student_id=sid, period_start=today, period_end=today, attendance_rate=rate,
                           total_sessions=number(row, "total_sessions"),
                           absent_sessions=number(row, "absent_sessions")),
                LearningInteraction(student_id=sid, period_start=today, period_end=today,
                                    login_count=number(row, "login_count"),
                                    assignment_submitted=number(row, "assignment_submitted"),
                                    assignment_missing=number(row, "assignment_missing"),
                                    forum_posts=number(row, "forum_posts"), video_views=number(row, "video_views"),
                                    learning_hours=number(row, "learning_hours", float)),
            ]
            imported_rows += 1

    return {"imported": imported_rows if _commit(db, pending, errors) else 0, "errors": errors}


def template_csv(kind: str) -> tuple[bytes, str]:
    samples = {
        "students": (STUDENT_COLUMNS, ["SV20250001", "Nguyễn Văn An", "an@example.edu.vn", "0901234567",
                                       "CNTT-K25-01", "Công nghệ thông tin", "2025", "active"],
                     "mau_sinh_vien.csv"),
        "metrics": (METRIC_COLUMNS, ["SV20250001", "2025.1", "7.5", "0", "68", "18", "86.5", "45", "6",
                                     "24", "11", "1", "3", "28", "12.5"], "mau_chi_so.csv"),
    }
    if kind not in samples:
        raise ImportFileError(f"Không có file mẫu '{kind}'")
    columns, sample, filename = samples[kind]
    # BOM để Excel trên Windows mở đúng tiếng Việt.
    return (",".join(columns) + "\n" + ",".join(sample) + "\n").encode("utf-8-sig"), filename
