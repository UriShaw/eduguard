"""
Xuất báo cáo Excel theo lớp và học kỳ.

Giới hạn dữ liệu cần nói rõ, vì nó ảnh hưởng tới cách đọc báo cáo: chỉ bảng
academic_results có cột học kỳ. Chuyên cần và mức nguy cơ in ra là bản ghi
GẦN NHẤT của sinh viên tại thời điểm xuất file, không phải số liệu riêng của
học kỳ đang lọc. Ghi chú này được in ngay trong file Excel để người nhận
không hiểu nhầm.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.extensions import db
from app.ml.features import AT_RISK_LEVELS
from app.models import AcademicResult, Student
from app.services import latest

PRIMARY = "004AC6"
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=PRIMARY)
TITLE_FONT = Font(name="Arial", size=14, bold=True, color=PRIMARY)
NOTE_FONT = Font(name="Arial", size=9, italic=True, color="666666")
BODY_FONT = Font(name="Arial", size=10)
BORDER = Border(*(Side(style="thin", color="D0D0D0") for _ in range(4)))

NOTE = ("Chuyên cần và mức nguy cơ là số liệu GẦN NHẤT của sinh viên tại thời điểm "
        "xuất báo cáo, không tách riêng theo từng học kỳ.")


def available_semesters(advisor_user_id: int | None = None) -> list[str]:
    query = db.session.query(AcademicResult.semester).distinct()
    if advisor_user_id is not None:
        query = query.join(Student, AcademicResult.student_id == Student.student_id).filter(
            Student.advisor_user_id == advisor_user_id
        )
    return sorted((row[0] for row in query.all()), reverse=True)


def collect_rows(class_name: str | None = None, semester: str | None = None,
                 advisor_user_id: int | None = None) -> list[dict]:
    """Mỗi dòng là một cặp (sinh viên, kết quả học kỳ) khớp bộ lọc."""
    query = (
        db.session.query(AcademicResult, Student)
        .join(Student, AcademicResult.student_id == Student.student_id)
    )
    if class_name:
        query = query.filter(Student.class_name == class_name)
    if semester:
        query = query.filter(AcademicResult.semester == semester)
    if advisor_user_id is not None:
        query = query.filter(Student.advisor_user_id == advisor_user_id)

    pairs = query.order_by(Student.class_name, AcademicResult.semester, Student.full_name).all()

    # Hai truy vấn cố định cho toàn bộ báo cáo, thay vì hai truy vấn mỗi dòng.
    student_ids = list({student.student_id for _, student in pairs})
    attendance = latest.attendance_by_student(student_ids)
    prediction = latest.prediction_by_student(student_ids)

    rows = []
    for record, student in pairs:
        att = attendance.get(student.student_id)
        pred = prediction.get(student.student_id)
        rows.append({
            "class_name": student.class_name or "—",
            "semester": record.semester,
            "student_code": student.student_code,
            "full_name": student.full_name,
            "gpa": float(record.gpa),
            "failed_subjects": record.failed_subjects,
            "attendance_rate": float(att.attendance_rate) if att else None,
            "risk_level": pred.risk_level if pred else "Chưa dự đoán",
        })
    return rows


def _write_header(sheet, row: int, headers: list[str]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def _write_row(sheet, row: int, values: list, center_from: int | None = None) -> None:
    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=col, value=value)
        cell.font = BODY_FONT
        cell.border = BORDER
        if center_from and col >= center_from:
            cell.alignment = Alignment(horizontal="center")


def _set_widths(sheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width


def _build_summary_sheet(sheet, rows: list[dict]) -> None:
    sheet.title = "Tổng hợp"
    sheet["A1"] = "BÁO CÁO NGUY CƠ BỎ HỌC THEO LỚP VÀ HỌC KỲ"
    sheet["A1"].font = TITLE_FONT
    sheet.merge_cells("A1:H1")
    sheet["A2"] = NOTE
    sheet["A2"].font = NOTE_FONT
    sheet.merge_cells("A2:H2")

    headers = ["Lớp", "Học kỳ", "Số SV", "GPA trung bình", "Chuyên cần TB (%)",
               "Nguy cơ thấp", "Nguy cơ trung bình", "Nguy cơ cao"]
    _write_header(sheet, 4, headers)

    groups: dict[tuple[str, str], list[dict]] = {}
    for row in rows:
        groups.setdefault((row["class_name"], row["semester"]), []).append(row)

    line = 5
    for (class_name, semester), items in sorted(groups.items()):
        gpas = [i["gpa"] for i in items]
        rates = [i["attendance_rate"] for i in items if i["attendance_rate"] is not None]

        _write_row(sheet, line, [
            class_name,
            semester,
            len(items),
            round(sum(gpas) / len(gpas), 2) if gpas else None,
            round(sum(rates) / len(rates), 1) if rates else None,
            sum(1 for i in items if i["risk_level"] == "Thấp"),
            sum(1 for i in items if i["risk_level"] == "Trung bình"),
            sum(1 for i in items if i["risk_level"] in AT_RISK_LEVELS),
        ], center_from=3)
        line += 1

    if not groups:
        sheet.cell(row=line, column=1, value="(Không có dữ liệu khớp bộ lọc)").font = NOTE_FONT

    _set_widths(sheet, [16, 12, 10, 16, 18, 14, 20, 14])


def _build_detail_sheet(sheet, rows: list[dict]) -> None:
    headers = ["Lớp", "Học kỳ", "Mã SV", "Họ và tên", "GPA", "Môn trượt",
               "Chuyên cần gần nhất (%)", "Mức nguy cơ gần nhất"]
    _write_header(sheet, 1, headers)

    for line, row in enumerate(rows, start=2):
        _write_row(sheet, line, [
            row["class_name"], row["semester"], row["student_code"], row["full_name"],
            row["gpa"], row["failed_subjects"], row["attendance_rate"], row["risk_level"],
        ])

    if not rows:
        sheet.cell(row=2, column=1, value="(Không có dữ liệu khớp bộ lọc)").font = NOTE_FONT

    _set_widths(sheet, [16, 12, 14, 26, 8, 12, 22, 22])
    sheet.freeze_panes = "A2"


def build_workbook(class_name: str | None = None, semester: str | None = None,
                   advisor_user_id: int | None = None) -> BytesIO:
    """File Excel hai sheet: Tổng hợp theo lớp+kỳ, và Chi tiết từng sinh viên."""
    rows = collect_rows(class_name, semester, advisor_user_id)

    workbook = Workbook()
    _build_summary_sheet(workbook.active, rows)
    _build_detail_sheet(workbook.create_sheet("Chi tiết"), rows)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def filename_for(class_name: str | None, semester: str | None) -> str:
    parts = ["bao_cao"]
    if class_name:
        parts.append(class_name.replace(" ", "_"))
    if semester:
        parts.append(semester.replace(" ", "_"))
    return "_".join(parts) + ".xlsx"
